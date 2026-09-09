"""Line-level surplus-lines parse + NAIC Companies row without a CoCode."""
from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "colorado" / "co-ins-001"
TEXT = (RAW / "surplus-extracted.txt").read_text(encoding="utf-8")
OUT = ROOT / "artifacts" / "co-ins-001" / "surplus-line-parse.json"

ROW_RE = re.compile(
    r"^(?P<name>.+?)\s+(?P<id>(?:\d{5}|AA-\d{7}))\s*$",
    re.I,
)


def main() -> None:
    rows = []
    unmatched = []
    for i, raw in enumerate(TEXT.splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue
        if line in {"1", "2", "3", "4"}:
            continue
        if line.lower().startswith("insurer name"):
            continue
        m = ROW_RE.match(line)
        if m:
            ident = m.group("id").upper()
            rows.append({"line": i, "name": m.group("name").strip(), "id": ident, "kind": "alien" if ident.startswith("AA-") else "naic"})
        else:
            unmatched.append({"line": i, "text": line})

    naic = [r for r in rows if r["kind"] == "naic"]
    alien = [r for r in rows if r["kind"] == "alien"]
    ids = [r["id"] for r in rows]
    dup = sorted({i for i in ids if ids.count(i) > 1})

    wb = load_workbook(RAW / "Colorado-2025-Statistical-Report.xlsx", read_only=True, data_only=True)
    ws = wb["NAIC Companies"]
    header = None
    missing = []
    all_rows = 0
    for rec in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip() if c is not None else "" for c in rec]
            continue
        if not rec or all(c is None or str(c).strip() == "" for c in rec):
            continue
        all_rows += 1
        raw_naic = rec[0]
        digits = "".join(ch for ch in str(raw_naic or "") if ch.isdigit())
        if not (3 <= len(digits) <= 5):
            missing.append({"naic_raw": raw_naic, "name": rec[1], "dom": rec[2]})
    wb.close()

    payload = {
        "surplus_line_rows": len(rows),
        "surplus_naic_rows": len(naic),
        "surplus_alien_rows": len(alien),
        "surplus_distinct_ids": len(set(ids)),
        "surplus_duplicate_ids": dup,
        "unmatched_lines": unmatched,
        "sample_rows": rows[:8],
        "tail_rows": rows[-8:],
        "naic_companies_data_rows": all_rows,
        "naic_companies_rows_without_cocode": missing,
    }
    OUT.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, indent=2)[:6000])


if __name__ == "__main__":
    main()
