"""CO-INS-001A2 census: surplus 2026-27, complaints, certs, NAIC crosswalk."""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import ssl
import urllib.request
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[2]
RAW = ROOT / "data" / "colorado" / "co-ins-001"
ART = ROOT / "artifacts" / "co-ins-001"
UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
CTX = ssl.create_default_context()
CTX.check_hostname = False
CTX.verify_mode = ssl.CERT_NONE
ROW_RE = re.compile(r"^(?P<name>.+?)\s+(?P<id>(?:\d{5}|AA-\d{7}))\s*$", re.I)
STD_LINES = [
    ("AUTO", "Automobile, Private Passenger"),
    ("HOME", "Homeowners"),
    ("ACHL", "Accident and Health"),
    ("ANNT", "Annuity"),
    ("LIFE", "Life"),
    ("HMO", "HMO"),
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def fetch(url: str) -> tuple[int, bytes]:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    try:
        with urllib.request.urlopen(req, timeout=90, context=CTX) as resp:
            return resp.status, resp.read()
    except Exception as e:
        return 0, str(e).encode()


def parse_surplus_pdf(path: Path) -> dict:
    reader = PdfReader(str(path))
    text = "\n".join((p.extract_text() or "") for p in reader.pages)
    (RAW / f"{path.stem}-extracted.txt").write_text(text, encoding="utf-8")
    rows = []
    unmatched = []
    for i, raw in enumerate(text.splitlines(), start=1):
        line = raw.strip()
        if not line or line in {"1", "2", "3", "4", "5"}:
            continue
        if line.lower().startswith("insurer name"):
            continue
        m = ROW_RE.match(line)
        if m:
            ident = m.group("id").upper()
            rows.append(
                {
                    "name": m.group("name").strip(),
                    "id": ident,
                    "kind": "alien" if ident.startswith("AA-") else "naic",
                }
            )
        else:
            unmatched.append(line)
    ids = [r["id"] for r in rows]
    naic = [r["id"] for r in rows if r["kind"] == "naic"]
    alien = [r["id"] for r in rows if r["kind"] == "alien"]
    return {
        "filename": path.name,
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "pages": len(reader.pages),
        "text_head": text[:1200],
        "eligible_identities": len(rows),
        "naic_cocode_identities": len(naic),
        "alien_aa_identities": len(alien),
        "distinct_ids": len(set(ids)),
        "duplicate_ids": sorted({i for i in ids if ids.count(i) > 1}),
        "effective_from": "2026-07-01" if "July 1, 2026" in text or "2026 through" in text else None,
        "effective_through": "2027-06-30" if "June 30, 2027" in text else None,
        "updated_at": "2026-07-24" if ("7/24/26" in text or "7/24/2026" in text or "July 24" in text) else None,
        "has_july_1_2026": "July 1, 2026" in text,
        "has_june_30_2027": "June 30, 2027" in text,
        "unmatched_sample": unmatched[:20],
        "naic_ids": sorted(set(naic)),
        "alien_ids": sorted(set(alien)),
    }


def parse_surplus_2025() -> dict:
    text_path = RAW / "surplus-extracted.txt"
    pdf = RAW / "colorado-eligible-nonadmitted-2025-2026.pdf"
    if not text_path.exists() and pdf.exists():
        reader = PdfReader(str(pdf))
        text = "\n".join((p.extract_text() or "") for p in reader.pages)
    else:
        text = text_path.read_text(encoding="utf-8") if text_path.exists() else ""
    rows = []
    for raw in text.splitlines():
        m = ROW_RE.match(raw.strip())
        if m:
            ident = m.group("id").upper()
            rows.append({"id": ident, "kind": "alien" if ident.startswith("AA-") else "naic", "name": m.group("name").strip()})
    naic = {r["id"] for r in rows if r["kind"] == "naic"}
    alien = {r["id"] for r in rows if r["kind"] == "alien"}
    return {
        "eligible_identities": len(rows),
        "naic_ids": sorted(naic),
        "alien_ids": sorted(alien),
    }


def statistical_naic() -> list[str]:
    path = RAW / "Colorado-2025-Statistical-Report.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["NAIC Companies"]
    codes: list[str] = []
    header = None
    for rec in ws.iter_rows(values_only=True):
        if header is None:
            header = [str(c).strip() if c is not None else "" for c in rec]
            continue
        if not rec or all(c is None or str(c).strip() == "" for c in rec):
            continue
        raw = rec[0]
        digits = "".join(ch for ch in str(raw or "") if ch.isdigit())
        if 3 <= len(digits) <= 5:
            codes.append(digits.zfill(5))
    wb.close()
    return codes


def parse_certs() -> dict:
    html = (RAW / "certs.html").read_text(encoding="utf-8", errors="replace")
    bodies = re.findall(
        r'field--name-field-card-body.*?field--item">(.*?)</div>',
        html,
        flags=re.I | re.S,
    )
    items = []
    for body in bodies:
        for m in re.finditer(r"<li[^>]*>(.*?)</li>", body, flags=re.I | re.S):
            chunk = m.group(1)
            text = re.sub(r"<[^>]+>", " ", chunk)
            text = re.sub(r"\s+", " ", text).replace("&amp;", "&").replace("&nbsp;", " ").strip()
            hrefs = re.findall(r'href="([^"]+)"', chunk)
            if text:
                items.append({"name": text, "hrefs": hrefs})
    heading = [i for i in items if i["name"].strip().lower() == "title agencies"]
    companies = [i for i in items if i["name"].strip().lower() != "title agencies"]
    all_hrefs = [h for i in companies for h in i["hrefs"]]
    pdf_hrefs = [h for h in all_hrefs if ".pdf" in h.lower()]
    drive_hrefs = [h for h in all_hrefs if "drive.google.com" in h.lower()]
    multi = [i for i in companies if len(i["hrefs"]) > 1]
    no_link = [i for i in companies if not i["hrefs"]]
    return {
        "list_items_including_heading": len(items),
        "excluded_heading_count": len(heading),
        "named_company_entries": len(companies),
        "document_links": len(all_hrefs),
        "unique_document_links": len(set(all_hrefs)),
        "doi_pdf_links": len(pdf_hrefs),
        "google_drive_links": len(drive_hrefs),
        "companies_with_multiple_links": [{"name": i["name"], "hrefs": i["hrefs"]} for i in multi],
        "companies_with_no_link": [i["name"] for i in no_link],
        "names": [i["name"] for i in companies],
        "reconciliation": (
            f"{len(companies)} named company entries after excluding the Title Agencies heading; "
            f"{len(all_hrefs)} document links ({len(pdf_hrefs)} DOI PDF + {len(drive_hrefs)} Google Drive). "
            "Company-entry count and document-link count are different grains."
        ),
    }


def fetch_std_reports() -> dict:
    rows = []
    by_line = []
    for code, label in STD_LINES:
        url = (
            "https://www.dora.state.co.us/pls/real/ins_comp_ratio_report.generate_report"
            "?p_report_id=Complaint%20Ratio%20and%20Complaint%20Index%20Search%20Results"
            f"&p_report_year=2025&p_policy_type={code}&p_min_complaints=5"
            "&p_min_marketshare=0.10&p_report_format=HTML&p_report_type=Standard%20Reports"
            "&p_order_by=company_name"
        )
        status, body = fetch(url)
        html = body.decode("latin-1", "replace") if status == 200 else ""
        if status == 200:
            (RAW / f"complaint-std-2025-{code.lower()}.html").write_bytes(body)
        naics = re.findall(r"p_naic_code=(\d{5})", html)
        names = re.findall(r"p_naic_code=\d{5}[^>]*>([^<]+)</a>", html)
        line_rows = []
        # Pair company detail anchors with NAIC
        for m in re.finditer(
            r'p_naic_code=(\d{5})[^>]*>([^<]+)</a>',
            html,
        ):
            naic, name = m.group(1), re.sub(r"\s+", " ", m.group(2)).strip()
            alias = None
            am = re.search(r"\(([^)]+)\)\s*$", name)
            if am and am.group(1) not in {"S.I.", "Inc.", "INC.", "LLC"}:
                # Parenthetical may be DBA/common name — not an identity key.
                alias = am.group(1)
            line_rows.append({"naic": naic, "name": name, "alias_parenthetical": alias, "line": code})
        by_line.append(
            {
                "line_code": code,
                "line_label": label,
                "http_status": status,
                "bytes": len(body),
                "sha256": hashlib.sha256(body).hexdigest() if status == 200 and body else None,
                "url": url,
                "rows": len(line_rows),
                "distinct_naic": len({r["naic"] for r in line_rows}),
            }
        )
        rows.extend(line_rows)
        print("std", code, status, len(line_rows), "naic", len(set(naics)))
    return {"year": "2025", "lines": by_line, "rows": rows}


def parse_complaint_pdf() -> dict:
    path = RAW / "complaint_fy2024_25.pdf"
    reader = PdfReader(str(path))
    # First ~15 pages typically have facts/figures.
    text = "\n".join((p.extract_text() or "") for p in reader.pages[:20])
    (RAW / "complaint_fy2024_25-head.txt").write_text(text, encoding="utf-8")
    money = re.findall(r"\$[\d,]+", text)
    ints = re.findall(r"\b[\d,]{3,}\b", text)
    return {
        "filename": path.name,
        "bytes": path.stat().st_size,
        "sha256": sha256(path),
        "pages": len(reader.pages),
        "period": "FY 2024-25 (July 2024 - June 2025)",
        "source_as_of": "2025-06-30",
        "published_at": "2025-11",
        "text_head": text[:4000],
        "money_hits": money[:30],
        "int_hits": ints[:40],
        "mentions_confirmed": "confirmed" in text.lower(),
        "mentions_recoveries": "recover" in text.lower(),
        "mentions_complaint": "complaint" in text.lower(),
    }


def crosswalk(source_naic: list[str], spine: set[str]) -> dict:
    src = sorted(set(source_naic))
    matches = [c for c in src if c in spine]
    unmatched = [c for c in src if c not in spine]
    invalid = [c for c in src if not re.fullmatch(r"\d{5}", c)]
    return {
        "CROSSWALK_STATUS": "EXECUTED",
        "SOURCE_DISTINCT_NAIC": len(src),
        "EXACT_EXISTING_LEGAL_INSURER_MATCHES": len(matches),
        "UNMATCHED_NAIC": len(unmatched),
        "INVALID_OR_UNRESOLVED_IDS": len(invalid),
        "unmatched_sample": unmatched[:25],
        "invalid_sample": invalid[:10],
        "read_only": True,
        "graph_write": False,
    }


def main() -> None:
    surplus_new = parse_surplus_pdf(RAW / "surplus_2026_2027.pdf")
    surplus_old = parse_surplus_2025()
    new_naic = set(surplus_new["naic_ids"])
    old_naic = set(surplus_old["naic_ids"])
    surplus_delta = {
        "added_naic": sorted(new_naic - old_naic),
        "removed_naic": sorted(old_naic - new_naic),
        "stable_naic": len(new_naic & old_naic),
        "added_alien": sorted(set(surplus_new["alien_ids"]) - set(surplus_old["alien_ids"])),
        "removed_alien": sorted(set(surplus_old["alien_ids"]) - set(surplus_new["alien_ids"])),
    }
    certs = parse_certs()
    std = fetch_std_reports()
    complaints_pdf = parse_complaint_pdf()
    stat_naic = statistical_naic()
    spine = {
        r["naic_cocode"].zfill(5)
        for r in json.loads((ROOT / "data/reports/ins-insurer-006-identity-index.json").read_text(encoding="utf-8"))["insurers"]
        if r.get("naic_cocode")
    }
    stat_x = crosswalk(stat_naic, spine)
    surplus_x = crosswalk(surplus_new["naic_ids"], spine)
    complaint_naics = [r["naic"] for r in std["rows"]]
    complaint_x = crosswalk(complaint_naics, spine)
    complaint_labels = [r["name"] for r in std["rows"]]
    aliases = [r for r in std["rows"] if r.get("alias_parenthetical")]
    public = {
        "ticket": "CO-INS-001A2",
        "retrieved_at": "2026-09-09",
        "surplus_2026_2027": {k: v for k, v in surplus_new.items() if k not in {"naic_ids", "alien_ids"}},
        "surplus_2026_2027_naic_ids": surplus_new["naic_ids"],
        "surplus_2026_2027_alien_ids": surplus_new["alien_ids"],
        "surplus_2025_2026": {
            "eligible_identities": surplus_old["eligible_identities"],
            "naic": len(surplus_old["naic_ids"]),
            "alien": len(surplus_old["alien_ids"]),
            "classification": "HISTORICAL_ELIGIBILITY_SNAPSHOT",
            "effective_from": "2025-07-01",
            "effective_through": "2026-06-30",
            "updated_at": "2026-01-26",
        },
        "surplus_delta": {
            "added_naic_count": len(surplus_delta["added_naic"]),
            "removed_naic_count": len(surplus_delta["removed_naic"]),
            "stable_naic": surplus_delta["stable_naic"],
            "added_naic_sample": surplus_delta["added_naic"][:20],
            "removed_naic_sample": surplus_delta["removed_naic"][:20],
            "added_alien": surplus_delta["added_alien"],
            "removed_alien": surplus_delta["removed_alien"],
        },
        "certificates": certs,
        "complaints_annual": complaints_pdf,
        "complaints_standard_2025": {
            "access": "ACQUIRED",
            "year": "2025",
            "inclusion_rule": "five or more complaints OR 0.1% or more of premium volume in a line",
            "ratio_uses": "all received complaints, not confirmed-only",
            "lines": std["lines"],
            "row_count": len(std["rows"]),
            "distinct_naic": len(set(complaint_naics)),
            "distinct_company_labels": len(set(complaint_labels)),
            "rows_with_parenthetical_alias": len(aliases),
            "has_naic": True,
            "identity_key": "p_naic_code on official detail link",
        },
        "complaints_interactive": {
            "access": "OPEN_SEARCH_ONLY",
            "url": "http://www.dora.state.co.us/pls/real/ins_comp_ratio_report.interactive_report_page?p_report_id=Complaint%20Ratio%20and%20Complaint%20Index%20Search%20Results&p_label=",
            "inclusion_rule": "any Colorado insurance premium written OR at least one complaint",
            "scraped": False,
        },
        "file_a_complaint": {
            "access": "PUBLIC_RESEARCH_PATH",
            "url": "https://doi.colorado.gov/for-consumers/file-a-complaint",
        },
        "crosswalk": {
            "spine_source": "data/reports/ins-insurer-006-identity-index.json",
            "spine_legal_insurers": len(spine),
            "statistical_report": stat_x,
            "surplus_lines_naic": surplus_x,
            "complaint_standard_2025": complaint_x,
            "alien_aa_not_crosswalked": True,
        },
        "complaint_rows_for_snapshot": [
            {
                "naic": r["naic"],
                "name": r["name"],
                "line": r["line"],
                "alias_parenthetical": r["alias_parenthetical"],
            }
            for r in std["rows"]
        ],
    }
    # Do not store huge unmatched lists in the committed snapshot census beyond samples.
    (ART / "a2-census.json").write_text(json.dumps(public, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({
        "surplus_new": surplus_new["eligible_identities"],
        "surplus_naic": surplus_new["naic_cocode_identities"],
        "surplus_alien": surplus_new["alien_aa_identities"],
        "effective_from": surplus_new["effective_from"],
        "effective_through": surplus_new["effective_through"],
        "updated": surplus_new["updated_at"],
        "delta_added": len(surplus_delta["added_naic"]),
        "delta_removed": len(surplus_delta["removed_naic"]),
        "certs_companies": certs["named_company_entries"],
        "certs_links": certs["document_links"],
        "std_rows": len(std["rows"]),
        "std_naic": len(set(complaint_naics)),
        "stat_x": {k: stat_x[k] for k in list(stat_x)[:5]},
        "surplus_x": {k: surplus_x[k] for k in list(surplus_x)[:5]},
        "complaint_x": {k: complaint_x[k] for k in list(complaint_x)[:5]},
        "complaint_pdf_pages": complaints_pdf["pages"],
        "complaint_head_len": len(complaints_pdf["text_head"]),
    }, indent=2))


if __name__ == "__main__":
    main()
